"""
Пакетное создание договоров из Excel.

accept()  — проверяет файл, сохраняет пачку: строки с ошибками → rejected, остальные → pending.
Воркер    — берёт pending по одной (SELECT ... FOR UPDATE SKIP LOCKED) и создаёт договор
            через LeaseContractService. Очередь живёт в Postgres, поэтому переживает перезапуск
            и безопасна при нескольких репликах.
errors_workbook() — исходный Excel, в котором оставлены только строки с ошибками
            и добавлен столбец «Комментарий». Исправленный файл можно загрузить повторно:
            уже созданные договоры не задублируются (идемпотентность по номеру).
"""

from __future__ import annotations

import asyncio
import logging
import re
import uuid
from copy import copy
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from io import BytesIO

import openpyxl
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select, text, update
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from db.minio import MinioStorage
from db.models import LeaseBatch, LeaseBatchItem
from helpers.lease_excel import CONTRACTS_SHEET, FIRST_DATA_ROW, parse_lease_workbook
from schemas import SubleaseContractData
from services.lease_contract import LeaseContractService

logger = logging.getLogger(__name__)

STATUS_REJECTED = "rejected"  # не прошла проверку при загрузке
STATUS_PENDING = "pending"
STATUS_PROCESSING = "processing"
STATUS_DONE = "done"
STATUS_ERROR = "error"  # упала при создании (Gotenberg, MinIO, конфликт номера…)

_XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_COMMENT_HEADER = "Комментарий"
_STALE_AFTER = timedelta(minutes=10)

_CLAIM_NEXT = text(f"""
    UPDATE lease_batch_items
       SET status = '{STATUS_PROCESSING}', attempts = attempts + 1, updated = now()
     WHERE id = (
            SELECT id FROM lease_batch_items
             WHERE status = '{STATUS_PENDING}'
             ORDER BY id
             FOR UPDATE SKIP LOCKED
             LIMIT 1)
 RETURNING id, batch_id, payload
""")


class BatchFileRejected(Exception):
    """Файл нельзя принять целиком (нет листов, ошибки на листе «Объект»)."""

    def __init__(self, errors: list[str]):
        super().__init__("; ".join(errors))
        self.errors = errors


@dataclass
class BatchSummary:
    batch_id: uuid.UUID
    category: str
    total: int
    accepted: int
    rejected: int


@dataclass
class BatchRowError:
    row_number: int
    status: str
    error: str | None


@dataclass
class BatchStatus:
    batch_id: uuid.UUID
    filename: str | None
    created: datetime
    total: int
    counts: dict[str, int]
    finished: bool
    errors: list[BatchRowError] = field(default_factory=list)


class LeaseBatchService:
    def __init__(
        self,
        storage: MinioStorage,
        session_factory: async_sessionmaker[AsyncSession],
        lease_service: LeaseContractService,
    ):
        self._storage = storage
        self._session_factory = session_factory
        self._lease_service = lease_service

    # ── API ───────────────────────────────────────────────────────────────────

    async def accept(self, content: bytes, filename: str, created_by: str | None) -> BatchSummary:
        """Тип договора (аренда/субаренда) — из поля «Тип договора» на листе «Объект»."""
        parsed = await asyncio.to_thread(parse_lease_workbook, content)
        if parsed.fatal:
            raise BatchFileRejected(parsed.fatal)

        batch_id = uuid.uuid4()
        safe_name = re.sub(r"[\\/]", "_", filename or "contracts.xlsx")
        source_path = f"batches/{batch_id}/{safe_name}"
        await asyncio.to_thread(self._storage.upload, source_path, BytesIO(content), _XLSX_TYPE)

        try:
            async with self._session_factory() as session, session.begin():
                session.add(LeaseBatch(
                    id=batch_id,
                    category=parsed.category,
                    filename=filename,
                    source_path=source_path,
                    total=len(parsed.rows),
                    created_by=created_by,
                ))
                await session.flush()
                session.add_all([
                    LeaseBatchItem(
                        batch_id=batch_id,
                        row_number=row.row_number,
                        payload=row.data.model_dump(mode="json") if row.data else None,
                        status=STATUS_PENDING if row.data else STATUS_REJECTED,
                        error="; ".join(row.errors) or None,
                    )
                    for row in parsed.rows
                ])
        except Exception:
            await asyncio.to_thread(self._storage.delete, source_path)
            raise

        accepted = sum(1 for row in parsed.rows if row.data)
        logger.info("Пачка %s принята: %d строк, %d в очереди", batch_id, len(parsed.rows), accepted)
        return BatchSummary(batch_id, parsed.category, len(parsed.rows), accepted, len(parsed.rows) - accepted)

    async def status(self, batch_id: uuid.UUID) -> BatchStatus | None:
        async with self._session_factory() as session:
            batch = await session.get(LeaseBatch, batch_id)
            if batch is None:
                return None
            counts = dict((await session.execute(
                select(LeaseBatchItem.status, func.count())
                .where(LeaseBatchItem.batch_id == batch_id)
                .group_by(LeaseBatchItem.status)
            )).all())
            errors = (await session.execute(
                select(LeaseBatchItem)
                .where(LeaseBatchItem.batch_id == batch_id,
                       LeaseBatchItem.status.in_([STATUS_REJECTED, STATUS_ERROR]))
                .order_by(LeaseBatchItem.row_number)
            )).scalars().all()

        return BatchStatus(
            batch_id=batch.id,
            filename=batch.filename,
            created=batch.created,
            total=batch.total,
            counts=counts,
            finished=not counts.get(STATUS_PENDING) and not counts.get(STATUS_PROCESSING),
            errors=[BatchRowError(e.row_number, e.status, e.error) for e in errors],
        )

    async def errors_workbook(self, batch_id: uuid.UUID) -> tuple[bytes, str] | None:
        """Исходный Excel: на листе «Договора» оставлены только строки с ошибками + столбец «Комментарий»."""
        async with self._session_factory() as session:
            batch = await session.get(LeaseBatch, batch_id)
            if batch is None:
                return None
            items = (await session.execute(
                select(LeaseBatchItem.row_number, LeaseBatchItem.error)
                .where(LeaseBatchItem.batch_id == batch_id,
                       LeaseBatchItem.status.in_([STATUS_REJECTED, STATUS_ERROR]))
            )).all()

        original = await asyncio.to_thread(self._storage.download, batch.source_path)
        content = await asyncio.to_thread(_build_errors_workbook, original, dict(items))
        name = f"Ошибки_{batch.filename or 'договоры.xlsx'}"
        return content, name

    async def retry(self, batch_id: uuid.UUID) -> int:
        """Строки, упавшие при создании (error), снова в очередь. rejected не трогаем — их надо исправить в файле."""
        async with self._session_factory() as session, session.begin():
            result = await session.execute(
                update(LeaseBatchItem)
                .where(LeaseBatchItem.batch_id == batch_id, LeaseBatchItem.status == STATUS_ERROR)
                .values(status=STATUS_PENDING, error=None, updated=func.now())
            )
            return result.rowcount

    # ── Воркер ────────────────────────────────────────────────────────────────

    async def reset_stale(self) -> int:
        """processing дольше _STALE_AFTER — воркер умер посреди строки (перезапуск контейнера)."""
        async with self._session_factory() as session, session.begin():
            result = await session.execute(
                update(LeaseBatchItem)
                .where(LeaseBatchItem.status == STATUS_PROCESSING,
                       LeaseBatchItem.updated < func.now() - _STALE_AFTER)
                .values(status=STATUS_PENDING, updated=func.now())
            )
            return result.rowcount

    async def process_next(self) -> bool:
        """Обрабатывает одну строку очереди. False — очередь пуста."""
        async with self._session_factory() as session, session.begin():
            claimed = (await session.execute(_CLAIM_NEXT)).first()
        if claimed is None:
            return False

        item_id, batch_id, payload = claimed
        async with self._session_factory() as session:
            batch = await session.get(LeaseBatch, batch_id)

        values: dict = {"updated": func.now()}
        try:
            result = await self._lease_service.create(
                SubleaseContractData.model_validate(payload),  # тип договора — внутри payload
                created_by=batch.created_by,
            )
            values.update(status=STATUS_DONE, error=None, lease_contract_id=result.id)
        except Exception as e:
            logger.exception("Строка %s пачки %s: ошибка создания договора", item_id, batch_id)
            values.update(status=STATUS_ERROR, error=str(e)[:2000] or type(e).__name__)

        async with self._session_factory() as session, session.begin():
            await session.execute(update(LeaseBatchItem).where(LeaseBatchItem.id == item_id).values(**values))
        return True


def _build_errors_workbook(original: bytes, comments: dict[int, str | None]) -> bytes:
    wb = openpyxl.load_workbook(BytesIO(original))
    ws = wb[CONTRACTS_SHEET]

    header_row = FIRST_DATA_ROW - 2  # строка 3 — заголовки колонок
    comment_col = ws.max_column + 1
    header = ws.cell(header_row, comment_col, _COMMENT_HEADER)
    template = ws.cell(header_row, comment_col - 1)
    header.font, header.fill, header.alignment, header.border = (
        copy(template.font), copy(template.fill), copy(template.alignment), copy(template.border)
    )
    ws.column_dimensions[get_column_letter(comment_col)].width = 70

    # Снизу вверх, чтобы удаление строк не сдвигало ещё не обработанные
    for row in range(ws.max_row, FIRST_DATA_ROW - 1, -1):
        if row in comments:
            ws.cell(row, comment_col, comments[row] or "Ошибка")
        elif any(ws.cell(row, c).value not in (None, "") for c in range(1, comment_col)):
            ws.delete_rows(row)

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


class LeaseBatchWorker:
    """Фоновая задача внутри приложения: обрабатывает очередь, пока приложение работает."""

    def __init__(self, service: LeaseBatchService, idle_seconds: float = 2.0):
        self._service = service
        self._idle_seconds = idle_seconds
        self._task: asyncio.Task | None = None

    def start(self) -> None:
        self._task = asyncio.create_task(self._run(), name="lease-batch-worker")

    async def stop(self) -> None:
        if self._task is not None:
            self._task.cancel()
            try:
                await self._task
            except asyncio.CancelledError:
                pass

    async def _run(self) -> None:
        try:
            reset = await self._service.reset_stale()
            if reset:
                logger.warning("Воркер: %d зависших строк возвращены в очередь", reset)
        except Exception:
            logger.exception("Воркер: не удалось вернуть зависшие строки")

        logger.info("Воркер пакетной загрузки запущен")
        while True:
            try:
                if not await self._service.process_next():
                    await asyncio.sleep(self._idle_seconds)
            except asyncio.CancelledError:
                raise
            except Exception:
                logger.exception("Воркер: ошибка, повтор через %s с", self._idle_seconds * 5)
                await asyncio.sleep(self._idle_seconds * 5)
